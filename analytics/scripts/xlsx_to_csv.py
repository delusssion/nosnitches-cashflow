"""Выгрузка данных экспертов в CSV по контракту.

Лист «Тестовые сделки» → testdata/deals/test_deals.csv (формат 1),
лист «КС ЦБ» → testdata/rates/key_rate.csv (формат 2).

Запуск: python analytics/scripts/xlsx_to_csv.py
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import sys
from collections import Counter

import openpyxl

from common import DEAL_COLUMNS, DEALS_CSV, KEY_RATE_CSV, SOURCE_XLSX

# В текстовых полях встречаются даты вида «2022-08-31 00:00:00» — время отрезаем.
MIDNIGHT = re.compile(r"(\d{4}-\d{2}-\d{2}) 00:00:00")


def fmt(value) -> str:
    """Значение ячейки → строка CSV по соглашениям контракта."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # Кратчайшая запись без хвостов двоичной арифметики: 0.075, 3646037.89, 156750.
        text = repr(round(value, 10))
        return text[:-2] if text.endswith(".0") else text
    text = str(value).strip()
    return MIDNIGHT.sub(r"\1", text)


def export_deals(wb) -> list[dict[str, str]]:
    ws = wb["Тестовые сделки"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    deals = []
    # rows[1] — описания полей, пропускаем; колонку «Столбец1» выбрасываем.
    for raw in rows[2:]:
        if all(v is None for v in raw):
            continue
        record = dict(zip(header, raw))
        deals.append({col: fmt(record.get(col)) for col in DEAL_COLUMNS})
    missing = [c for c in DEAL_COLUMNS if c not in header]
    if missing:
        sys.exit(f"В листе нет колонок: {missing}")
    DEALS_CSV.parent.mkdir(parents=True, exist_ok=True)
    with DEALS_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=DEAL_COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows(deals)
    return deals


def export_key_rate(wb) -> list[tuple[str, str]]:
    ws = wb["КС ЦБ"]
    rates = []
    # Сверху шапка-источник; данные — строки, где в первой колонке дата.
    for date, rate in ws.iter_rows(values_only=True, max_col=2):
        if isinstance(date, dt.datetime) and isinstance(rate, (int, float)):
            rates.append((date.date().isoformat(), fmt(float(rate))))
    rates.sort()
    KEY_RATE_CSV.parent.mkdir(parents=True, exist_ok=True)
    with KEY_RATE_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["indicator", "date", "rate"])
        w.writerows(["CBKEY1M", d, r] for d, r in rates)
    return rates


def main() -> None:
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True)
    deals = export_deals(wb)
    rates = export_key_rate(wb)
    print(f"Сделок: {len(deals)} → {DEALS_CSV.relative_to(DEALS_CSV.parents[2])}")
    for field in ("discriminator", "rate_type_cd", "day_convention_cd", "currency_id"):
        counts = Counter(d[field] for d in deals)
        print(f"  {field}: " + ", ".join(f"{k} {v}" for k, v in counts.most_common()))
    print(f"Ключевая ставка: {len(rates)} значений, {rates[0][0]} … {rates[-1][0]}")


if __name__ == "__main__":
    main()

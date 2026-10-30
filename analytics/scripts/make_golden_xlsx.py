"""Эталон в Excel с живыми формулами: analytics/golden_check.xlsx.

По листу на ключевые сделки (1, 2, 16, 22, 47). Параметры сделки — ячейки сверху, в графике —
ФОРМУЛЫ: даты (по правилу периода и номеру периода k), дни, доля года, начисление, проценты,
тело, остаток, субсидия. Каждое число можно проверить кликом.

Какие строки есть в графике, берётся из эталона cf_reference (даты будущих событий); сами
значения считает Excel / LibreOffice. Проверка: tests/test_golden_xlsx.py пересчитывает файл
LibreOffice в headless-режиме и сравнивает с testdata/golden (допуск 0.01).

Запуск: python analytics/scripts/make_golden_xlsx.py
"""

from __future__ import annotations

import datetime as dt
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "reference"))
from cf_reference import (  # noqa: E402
    Indicator,
    explicit_amounts,
    generate,
    interest_dates_all,
    is_future,
    load_deals,
    load_key_rate,
    make_client_leg,
    make_sub_leg,
    parse_date,
    parse_int,
    shift_by_lag,
)
from common import DEALS_CSV, KEY_RATE_CSV, ROOT  # noqa: E402

OUT = ROOT / "analytics" / "golden_check.xlsx"
XLSX_DEALS = [1, 2, 16, 22, 47]

HEAD = Font(bold=True)
PARAM_FILL = PatternFill("solid", fgColor="D5EBE4")
DATE_FMT = "DD.MM.YYYY"
MONEY_FMT = "#,##0.00"
RATE_FMT = "0.000%"

# Колонки графика.
COLS = [
    "k", "Дата", "Что происходит", "Дней", "Доля года τ", "Остаток на отрезке B", "Ставка r на отрезке",
    "Начислено B·r·τ", "Проценты к выплате", "Тело", "Остаток после",
    "Ставка субсидии на отрезке", "Субсидия начислено", "Субсидия к выплате",
]
C = {name: get_column_letter(i + 1) for i, name in enumerate(COLS)}
N = {name: i + 1 for i, name in enumerate(COLS)}
FIRST_ROW = 22  # строка «начало периода»; события — ниже


@dataclass
class Rule:
    """Правило дат «начало + N месяцев, число day» с ячейками параметров на листе."""

    start: dt.date
    months: int
    day: int | None  # None — число из даты начала, -1 — последний день
    start_cell: str
    day_cell: str
    months_cell: str

    def k_of(self, d: dt.date) -> int | None:
        k = (d.year - self.start.year) * 12 + d.month - self.start.month
        if k % self.months:
            return None
        return k // self.months

    def formula(self, k_cell: str, maturity_cell: str) -> str:
        """Дата периода k: число day в месяце start + k·N; если такого дня нет — последний день месяца."""
        s, n, day = self.start_cell, self.months_cell, self.day_cell
        first = f"DATE(YEAR({s}),MONTH({s})+{k_cell}*{n},1)"
        return f"=MIN(DATE(YEAR({s}),MONTH({s})+{k_cell}*{n},MIN({day},DAY(EOMONTH({first},0)))),{maturity_cell})"


def isda_formula(a: str, b: str) -> str:
    """ACT/ACT ISDA между датами a и b (ячейки): дни каждого календарного года / число дней в нём."""
    dy = lambda y: f"(DATE({y}+1,1,1)-DATE({y},1,1))"  # noqa: E731
    return (
        f"=IF(YEAR({a})=YEAR({b}),({b}-{a})/{dy(f'YEAR({a})')},"
        f"(DATE(YEAR({a})+1,1,1)-{a})/{dy(f'YEAR({a})')}+({b}-DATE(YEAR({b}),1,1))/{dy(f'YEAR({b})')}"
        f"+YEAR({b})-YEAR({a})-1)"
    )


def build_sheet(wb: Workbook, deal: dict[str, str], key_rate) -> None:
    i = int(deal["deal_id"])
    ws = wb.create_sheet(f"Сделка {i}")
    report = parse_date(deal["report_dttm"])
    maturity = parse_date(deal["maturity_dttm"])
    deal_date = parse_date(deal["deal_dttm"])
    disc = deal["discriminator"]
    events = generate(deal, key_rate, report)

    # --- Параметры -------------------------------------------------------------------------
    params: list[tuple[str, object, str | None]] = [
        ("Сделка", i, None),
        ("Тип погашения", disc, None),
        ("Дата расчёта T0", report, DATE_FMT),
        ("Дата погашения Tm", maturity, DATE_FMT),
        ("Остаток на T0, B0", float(deal["current_principal_amt"]), MONEY_FMT),
        ("Конвенция", deal["day_convention_cd"], None),
        ("Тип ставки", deal["rate_type_cd"], None),
        ("interest_rate (ставка / маржа)", float(deal["interest_rate"]), RATE_FMT),
    ]
    leg = make_client_leg(deal, deal_date, maturity)
    sub = make_sub_leg(deal, deal_date, maturity)

    # Строки графика: начало периода, прошлые пересмотры внутри периода, будущие события.
    by_date: dict[dt.date, dict[str, float]] = {}
    for e in events:
        if e["payment_type"] in ("repricing", "sub_repricing"):
            continue  # будущие пересмотры в этих сделках ставку не меняют (прогноз H3) — строки не нужны
        by_date.setdefault(e["payment_dttm"], {})[e["payment_type"]] = e["amount"]
    interest_all = interest_dates_all(deal, maturity)
    start = max(d for d in interest_all if not is_future(d, report))
    ind = Indicator(key_rate, report)
    repricings: dict[dt.date, str] = {}
    for leg_, label in ((leg, "пересмотр ставки"), (sub, "пересмотр ставки субсидии")):
        if leg_ and leg_.kind == "FLOAT":
            for d in leg_.repricing:
                prev_day = d - dt.timedelta(days=1)
                if start < d <= report and leg_.at(d, ind)[0] != leg_.at(prev_day, ind)[0]:
                    repricings[d] = label
    points = [start] + sorted(set(by_date) | set(repricings))

    def ks_key(leg_, t: dt.date) -> tuple[str, float]:
        """Параметр «КС ЦБ на дату фиксинга» для ставки, действующей с момента t."""
        fix = shift_by_lag(leg_.fixing_date(t), leg_.lag)
        if fix > report:
            return "КС ЦБ: прогноз после T0 (H3: значение на T0, сдвиг 0)", ind.value(fix)
        return f"КС ЦБ на {fix.strftime('%d.%m.%Y')} (фиксинг)", ind.value(fix)

    ks_params: dict[str, float] = {}
    for leg_ in (leg, sub):
        if leg_ and leg_.kind == "FLOAT":
            for t in points[:-1]:
                key, value = ks_key(leg_, t)
                ks_params[key] = value
    if leg.kind == "FLOAT":
        params += [
            ("floor", leg.floor if leg.floor is not None else 0, RATE_FMT),
            ("cap", leg.cap if leg.cap is not None else 1, RATE_FMT),
        ]
    if sub:
        params += [
            ("Субсидия: sub_interest_rate s", sub.spread, RATE_FMT),
            ("Субсидия: коэффициент k (пусто → 1)", sub.coeff, None),
        ]
    params += [(k, v, RATE_FMT) for k, v in sorted(ks_params.items())]

    ws["A1"] = f"Сделка {i}: эталон с формулами (методика docs/METHODOLOGY.md)"
    ws["A1"].font = Font(bold=True, size=13)
    cell: dict[str, str] = {}
    row = 3
    for label, value, fmt in params:
        ws.cell(row, 1, label)
        c = ws.cell(row, 2, value)
        c.fill = PARAM_FILL
        if fmt:
            c.number_format = fmt
        cell[label] = f"$B${row}"
        row += 1

    def client_rate(t: dt.date) -> str:
        """Формула ставки клиента, действующей с момента t: min(cap, max(floor, КС + маржа)) или FIXED."""
        if leg.kind == "FLOAT":
            ks = cell[ks_key(leg, t)[0]]
            return f"MIN({cell['cap']},MAX({cell['floor']},{ks}+{cell['interest_rate (ставка / маржа)']}))"
        return cell["interest_rate (ставка / маржа)"]

    def sub_rate(t: dt.date) -> str:
        """Ставка субсидии (H4): k · КС + s для FLOAT, s для FIXED."""
        if sub.kind == "FLOAT":
            ks = cell[ks_key(sub, t)[0]]
            return f"{cell['Субсидия: коэффициент k (пусто → 1)']}*{ks}+{cell['Субсидия: sub_interest_rate s']}"
        return cell["Субсидия: sub_interest_rate s"]

    ws.cell(row, 1, "Ставка клиента на T0 (для аннуитета)").font = HEAD
    ws.cell(row, 2, f"={client_rate(report)}").number_format = RATE_FMT
    rate_cell = f"$B${row}"
    row += 1

    # Правила дат (в столбцах D–E).
    rules: dict[str, Rule] = {}
    rrow = 3
    ws.cell(2, 4, "Правила дат").font = HEAD

    def add_rule(name: str, prefix: str) -> None:
        nonlocal rrow
        start = parse_date(deal[f"{prefix}_start_dttm"])
        period = deal[f"{prefix}_period_cd"]
        if not start or not period.endswith("M"):
            return
        day = parse_int(deal[f"{prefix}_day_num"])
        ws.cell(rrow, 4, f"{name}: начало")
        ws.cell(rrow, 5, start).number_format = DATE_FMT
        ws.cell(rrow + 1, 4, f"{name}: период, мес.")
        ws.cell(rrow + 1, 5, int(period[:-1]))
        ws.cell(rrow + 2, 4, f"{name}: число месяца")
        ws.cell(rrow + 2, 5, 31 if day == -1 else (day if day else start.day))
        for r in range(rrow, rrow + 3):
            ws.cell(r, 5).fill = PARAM_FILL
        rules[name] = Rule(start, int(period[:-1]), day, f"$E${rrow}", f"$E${rrow + 2}", f"$E${rrow + 1}")
        rrow += 3

    if deal["interest_payment_rule_period_cd"]:
        add_rule("Проценты", "interest_payment_rule")
    if deal["regular_payment_rule_period_cd"] not in ("", "ENDDATE"):
        add_rule("Тело", "regular_payment_rule")
    if sub:
        add_rule("Субсидия", "sub_payment_rule")
    if "пересмотр ставки" in repricings.values():
        add_rule("Пересмотр", "repricing_rule")
    if "пересмотр ставки субсидии" in repricings.values():
        add_rule("Пересмотр субсидии", "sub_repricing_rule")

    explicit = explicit_amounts(deal["principal_payment"]) if deal["principal_payment"] else {}
    exp_cell: dict[dt.date, tuple[str, str]] = {}
    if explicit:
        ws.cell(rrow, 4, "Явный график тела").font = HEAD
        rrow += 1
        for d, a in explicit.items():
            ws.cell(rrow, 4, d).number_format = DATE_FMT
            ws.cell(rrow, 5, a).number_format = MONEY_FMT
            ws.cell(rrow, 4).fill = ws.cell(rrow, 5).fill = PARAM_FILL
            exp_cell[d] = (f"$D${rrow}", f"$E${rrow}")
            rrow += 1

    # --- Строки графика ----------------------------------------------------------------------
    ws.cell(row + 1, 1, "Число будущих дат тела n").font = HEAD
    ws.cell(row + 1, 2, f'=COUNTIF({C["Что происходит"]}{FIRST_ROW + 1}:{C["Что происходит"]}{FIRST_ROW + len(points)},"*тело*")')
    n_cell = f"$B${row + 1}"
    ann_cell = None
    if disc == "ANN":
        ws.cell(row + 2, 1, "Аннуитет A = B0·i/(1−(1+i)^−n), i = r/12").font = HEAD
        ws.cell(row + 2, 2, f"={cell['Остаток на T0, B0']}*({rate_cell}/12)/(1-(1+{rate_cell}/12)^(-{n_cell}))")
        ws.cell(row + 2, 2).number_format = MONEY_FMT
        ann_cell = f"$B${row + 2}"
    if FIRST_ROW - 2 <= row + 2:
        raise RuntimeError("параметры не помещаются над графиком")

    for j, name in enumerate(COLS, start=1):
        c = ws.cell(FIRST_ROW - 1, j, name)
        c.font = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")

    last_interest_row = FIRST_ROW
    last_sub_row = FIRST_ROW
    for idx, d in enumerate(points):
        r = FIRST_ROW + idx
        kinds = by_date.get(d, {}) if idx else {}
        what = []
        if "interest_payment" in kinds:
            what.append("проценты")
        if "sub_payment" in kinds:
            what.append("субсидия")
        if "principal_payment" in kinds:
            what.append("тело")
        if "close_payment" in kinds:
            what.append("тело (закрытие)")
        if d in repricings:
            what.append(repricings[d])
        ws.cell(r, 3, "начало периода начисления" if idx == 0 else ", ".join(what))

        # Дата — формулой по правилу; явные даты — ссылкой на параметр; погашение — ссылкой на Tm.
        date_formula = None
        if d == maturity and idx:
            date_formula = f"={cell['Дата погашения Tm']}"
        elif d in exp_cell and "тело" in " ".join(what):
            date_formula = f"={exp_cell[d][0]}"
        else:
            first = ["Тело"] if "тело" in " ".join(what) else []
            first += ["Пересмотр субсидии"] if "пересмотр ставки субсидии" in what else []
            first += ["Пересмотр"] if "пересмотр ставки" in what else []
            for rule_name in first + ["Проценты", "Тело", "Субсидия"]:
                rule = rules.get(rule_name)
                k = rule.k_of(d) if rule else None
                if rule and k is not None and rule_date(rule, k, maturity) == d:
                    ws.cell(r, 1, k)
                    date_formula = rule.formula(f"{C['k']}{r}", cell["Дата погашения Tm"])
                    break
        if date_formula is None:
            raise RuntimeError(f"сделка {i}: не нашлось правила для даты {d}")
        ws.cell(r, 2, date_formula).number_format = DATE_FMT

        if idx == 0:
            ws.cell(r, N["Остаток после"], f"={cell['Остаток на T0, B0']}").number_format = MONEY_FMT
            continue

        prev = r - 1
        ws.cell(r, N["Дней"], f"={C['Дата']}{r}-{C['Дата']}{prev}")
        a, b = f"{C['Дата']}{prev}", f"{C['Дата']}{r}"
        tau = isda_formula(a, b) if deal["day_convention_cd"] in ("ACT/ACT ISDA", "ACT/ACT") else f"=({b}-{a})/365"
        ws.cell(r, N["Доля года τ"], tau).number_format = "0.0000000000"
        ws.cell(r, N["Остаток на отрезке B"], f"={C['Остаток после']}{prev}").number_format = MONEY_FMT
        # ставка на отрезке [prev, r) — та, что действует с даты prev
        ws.cell(r, N["Ставка r на отрезке"], f"={client_rate(points[idx - 1])}").number_format = RATE_FMT
        ws.cell(r, N["Начислено B·r·τ"],
                f"={C['Остаток на отрезке B']}{r}*{C['Ставка r на отрезке']}{r}*{C['Доля года τ']}{r}").number_format = MONEY_FMT
        if "interest_payment" in kinds:
            ws.cell(r, N["Проценты к выплате"], f"=SUM({C['Начислено B·r·τ']}{last_interest_row + 1}:{C['Начислено B·r·τ']}{r})")
            last_interest_row = r
        else:
            ws.cell(r, N["Проценты к выплате"], "=0")
        ws.cell(r, N["Проценты к выплате"]).number_format = MONEY_FMT

        bal = f"{C['Остаток на отрезке B']}{r}"
        if "close_payment" in kinds:
            body = f"={bal}"
        elif "principal_payment" in kinds:
            if disc == "ANN":
                body = f"={ann_cell}-{C['Проценты к выплате']}{r}"
            elif disc == "RGM":
                body = f"={cell['Остаток на T0, B0']}/{n_cell}"
            else:
                body = f"=MIN({exp_cell[d][1]},{bal})"
        else:
            body = "=0"
        ws.cell(r, N["Тело"], body).number_format = MONEY_FMT
        ws.cell(r, N["Остаток после"], f"={bal}-{C['Тело']}{r}").number_format = MONEY_FMT

        if sub:
            ws.cell(r, N["Ставка субсидии на отрезке"], f"={sub_rate(points[idx - 1])}").number_format = RATE_FMT
            ws.cell(r, N["Субсидия начислено"],
                    f"={bal}*{C['Ставка субсидии на отрезке']}{r}*{C['Доля года τ']}{r}").number_format = MONEY_FMT
            if "sub_payment" in kinds:
                ws.cell(r, N["Субсидия к выплате"], f"=SUM({C['Субсидия начислено']}{last_sub_row + 1}:{C['Субсидия начислено']}{r})")
                last_sub_row = r
            else:
                ws.cell(r, N["Субсидия к выплате"], "=0")
            ws.cell(r, N["Субсидия к выплате"]).number_format = MONEY_FMT

    end = FIRST_ROW + len(points) - 1
    tr = end + 2
    ws.cell(tr, 3, "Итого").font = HEAD
    for name in ("Проценты к выплате", "Тело") + (("Субсидия к выплате",) if sub else ()):
        col = C[name]
        ws[f"{col}{tr}"] = f"=SUM({col}{FIRST_ROW + 1}:{col}{end})"
        ws[f"{col}{tr}"].number_format = MONEY_FMT
        ws[f"{col}{tr}"].font = HEAD
    ws.cell(tr + 1, 3, "Проверка: тело = B0")
    ws.cell(tr + 1, N["Тело"], f"=ROUND({C['Тело']}{tr}-{cell['Остаток на T0, B0']},2)")
    ws.cell(tr + 2, 3, "Проверка: остаток после последнего = 0")
    ws.cell(tr + 2, N["Остаток после"], f"={C['Остаток после']}{end}")

    widths = [6, 12, 26, 7, 13, 17, 11, 16, 17, 16, 17, 11, 16, 16]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["D"].width = 24
    ws.freeze_panes = f"A{FIRST_ROW}"


def rule_date(rule: Rule, k: int, maturity: dt.date) -> dt.date:
    """Python-двойник Rule.formula — чтобы выбрать правило, дающее нужную дату."""
    import calendar

    m0 = rule.start.month - 1 + k * rule.months
    y, m = rule.start.year + m0 // 12, m0 % 12 + 1
    want = 31 if rule.day == -1 else (rule.day or rule.start.day)
    return min(dt.date(y, m, min(want, calendar.monthrange(y, m)[1])), maturity)


def main() -> None:
    deals = {int(d["deal_id"]): d for d in load_deals(DEALS_CSV)}
    key_rate = load_key_rate(KEY_RATE_CSV)
    wb = Workbook()
    readme = wb.active
    readme.title = "Как читать"
    lines = [
        "Эталонный расчёт с живыми формулами — nosnitches-cashflow",
        "",
        "Каждый лист — одна сделка. Зелёные ячейки — параметры сделки (вход), всё остальное — формулы.",
        "Даты: формула по правилу периода (начало, период, число месяца) и номеру периода k в колонке A;",
        "  явные даты графика — ссылка на параметр; последняя дата — дата погашения.",
        "Доля года: ACT/ACT ISDA — дни каждого календарного года / число дней в этом году; ACT/365 — дни / 365.",
        "Ставка на отрезке: FLOAT — min(cap, max(floor, КС на дату фиксинга + маржа)); значения КС — параметры;",
        "  после даты расчёта КС = значение на дату расчёта (гипотеза H3).",
        "Начислено на отрезке: B · r · τ; проценты к выплате — сумма начислений с прошлой выплаты.",
        "Тело: ANN — A − проценты; RGM — B0 / n; SIMPLE — по графику, не больше остатка; последнее — весь остаток.",
        "Первая строка — начало периода начисления (последняя выплата до даты расчёта): проценты до даты",
        "  расчёта входят в первый будущий платёж.",
        "Сверка с testdata/golden/deal_<id>.csv: analytics/tests/test_golden_xlsx.py (LibreOffice headless, допуск 0,01).",
        "Методика: docs/METHODOLOGY.md.",
    ]
    for r, text in enumerate(lines, start=1):
        readme.cell(r, 1, text)
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 110
    for i in XLSX_DEALS:
        build_sheet(wb, deals[i], key_rate)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: листы {', '.join(ws.title for ws in wb.worksheets)}")


if __name__ == "__main__":
    main()

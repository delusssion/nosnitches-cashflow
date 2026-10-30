"""Пересчёт листа «1 сделка» расчётного файла экспертов и проверка найденных ошибок.

Запуск: python analytics/scripts/check_expert_sheet.py
Печатает числа для docs/EXCEL_ERRORS.md.
"""

from __future__ import annotations

import datetime as dt
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "reference"))
from common import SOURCE_XLSX  # noqa: E402
from daycount import act_365, act_act_isda  # noqa: E402

FIRST_ROW, LAST_ROW = 19, 352  # строки графика на листе


@dataclass
class Sheet:
    principal: float
    rate: float
    n: int
    annuity: float
    last_paid: dt.date
    dates: list[dt.date]
    interest: list[float]
    principal_paid: list[float]
    balance_after_last: float
    formulas: dict[str, object]


def read_sheet(path: Path = SOURCE_XLSX) -> Sheet:
    wf = openpyxl.load_workbook(path)["1 сделка"]
    wv = openpyxl.load_workbook(path, data_only=True)["1 сделка"]
    rows = range(FIRST_ROW, LAST_ROW + 1)
    return Sheet(
        principal=wv["F3"].value,
        rate=wv["J3"].value,
        n=wv["B15"].value,
        annuity=wv["B14"].value,
        last_paid=wv["B17"].value.date(),
        dates=[wv[f"B{r}"].value.date() for r in rows],
        interest=[wv[f"F{r}"].value for r in rows],
        principal_paid=[wv[f"D{r}"].value for r in rows],
        balance_after_last=wv[f"E{LAST_ROW}"].value,
        formulas={c: wf[c].value for c in ("B14", "B15", "B19", "F19", "I19", "I20", f"F{LAST_ROW}")}
        | {"I": [wf[f"I{r}"].value for r in rows], "B": [wf[f"B{r}"].value for r in rows]},
    )


def annuity(principal: float, annual_rate: float, n: int) -> float:
    """Аннуитетный платёж: $A = B \\cdot i / (1 - (1+i)^{-n})$, $i = r/12$."""
    i = annual_rate / 12
    return principal * i / (1 - (1 + i) ** -n)


def schedule(sheet: Sheet, year_frac, close: bool) -> tuple[float, float]:
    """Аннуитетный график по датам листа. Возвращает (сумма процентов, остаток после последнего платежа)."""
    a = annuity(sheet.principal, sheet.rate, sheet.n)
    balance = sheet.principal
    prev = sheet.last_paid
    total = 0.0
    for k, d in enumerate(sheet.dates):
        interest = balance * sheet.rate * year_frac(prev, d)
        total += interest
        last = k == len(sheet.dates) - 1
        balance -= balance if (close and last) else a - interest
        prev = d
    return total, balance


def portfolio_refs(path: Path = SOURCE_XLSX) -> tuple[int, int]:
    ws = openpyxl.load_workbook(path)["Портфель 5 сделок"]
    deals = [row for row in ws.iter_rows(min_row=2, values_only=True) if isinstance(row[0], int)]
    broken = [row for row in deals if any("#REF!" in str(v) for v in row[1:])]
    return len(broken), len(deals)


def main() -> None:
    s = read_sheet()
    print("Лист «1 сделка», сделка 1: ANN, 3 646 037.89, 7.5 %, ACT/ACT, дата расчёта 2022-08-31")
    print(f"  платёж B14 = {s.annuity:.6f}; пересчёт A = {annuity(s.principal, s.rate, s.n):.6f}")
    print(f"  платежей B15 = {s.n} (константа: {s.formulas['B15']!r}); дат на листе: {len(s.dates)}")

    print("\n1. Аннуитет по r/12, проценты по дням /365")
    print(f"  остаток после последнего платежа (E{LAST_ROW}) = {s.balance_after_last:.2f}")
    t365, b365 = schedule(s, act_365, close=False)
    print(f"  пересчёт: остаток {b365:.2f}, проценты {t365:.2f}; сумма F на листе {sum(s.interest):.2f}")

    print("\n2. Конвенция ACT/ACT, а в формулах /365")
    tisda, bisda = schedule(s, act_act_isda, close=False)
    tisda_c, bisda_c = schedule(s, act_act_isda, close=True)
    t365_c, _ = schedule(s, act_365, close=True)
    print(f"  /365 + close_payment: проценты {t365_c:.2f}")
    print(f"  ISDA без закрытия: проценты {tisda:.2f}, остаток {bisda:.2f}")
    print(f"  ISDA + close_payment: проценты {tisda_c:.2f}, остаток {bisda_c:.2f}")
    print(f"  разница /365 − ISDA = {sum(s.interest) - tisda_c:.2f}")

    print("\n3. Ручной ввод")
    consts = sum(1 for v in s.formulas["B"] if not str(v).startswith("="))
    refs = sum(1 for v in s.formulas["I"] if str(v).startswith("=$") and "IF" not in str(v).upper())
    print(f"  дат платежей константами: {consts} из {len(s.dates)}; бакетов прямой ссылкой на заголовок: {refs}")
    print(f"  B15 = {s.formulas['B15']!r}, I19 = {s.formulas['I19']!r}, I20 = {s.formulas['I20']!r}")

    broken, total = portfolio_refs()
    print(f"\n4. Лист «Портфель 5 сделок»: #REF! у {broken} из {total} сделок")

    print("\nЧто сделано правильно")
    first_days = (s.dates[0] - s.last_paid).days
    print(f"  первый платёж процентов F19 = {s.interest[0]:.2f} = 3 646 037.89 × 7.5 % × {first_days}/365 — с {s.last_paid} (включая проценты до даты расчёта)")
    print(f"  аннуитет от остатка на дату расчёта: B14 = PMT(r/12, {s.n}, {s.principal})")


if __name__ == "__main__":
    main()

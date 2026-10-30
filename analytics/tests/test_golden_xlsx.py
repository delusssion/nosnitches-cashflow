"""golden_check.xlsx: формулы, пересчитанные LibreOffice, совпадают с testdata/golden (допуск 0.01)."""

import csv
import datetime as dt
import os
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from common import GOLDEN_DIR, ROOT
from make_golden_xlsx import COLS, FIRST_ROW, OUT, XLSX_DEALS

TOL = 0.01
COL = {name: i for i, name in enumerate(COLS)}
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice") or next(
    (p for p in ["/Applications/LibreOffice.app/Contents/MacOS/soffice"] if os.path.exists(p)), None
)


@pytest.fixture(scope="module")
def recalculated(tmp_path_factory) -> openpyxl.Workbook:
    if SOFFICE is None:
        pytest.skip("LibreOffice не установлен — пересчёт формул проверяется в CI")
    out = tmp_path_factory.mktemp("lo")
    profile = out / "profile"
    subprocess.run(
        [SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
         "--convert-to", "xlsx", "--outdir", str(out), str(OUT)],
        check=True, capture_output=True, timeout=300,
    )
    return openpyxl.load_workbook(out / OUT.name, data_only=True)


def golden(i: int) -> list[dict]:
    with (GOLDEN_DIR / f"deal_{i}.csv").open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def sheet_rows(ws):
    r = FIRST_ROW + 1
    while ws.cell(r, 2).value is not None and ws.cell(r, 3).value:
        yield [ws.cell(r, j + 1).value for j in range(len(COLS))]
        r += 1


@pytest.mark.parametrize("i", XLSX_DEALS)
def test_formulas_match_golden(recalculated, i):
    ws = recalculated[f"Сделка {i}"]
    by_date: dict[str, dict[str, dict]] = {}
    for e in golden(i):
        by_date.setdefault(e["payment_dttm"], {})[e["payment_type"]] = e
    rows = list(sheet_rows(ws))
    assert rows, "пустой график"
    checked = 0
    for row in rows:
        d = row[COL["Дата"]]
        d = (d.date() if isinstance(d, dt.datetime) else d).isoformat()
        if d not in by_date:
            # прошлый пересмотр ставки внутри первого периода: в выходе его нет, он только режет отрезок
            assert row[COL["Что происходит"]].startswith("пересмотр"), d
            continue
        ev = by_date.pop(d)
        if "interest_payment" in ev:
            assert row[COL["Проценты к выплате"]] == pytest.approx(float(ev["interest_payment"]["amount"]), abs=TOL), d
            checked += 1
        body = ev.get("principal_payment") or ev.get("close_payment")
        if body:
            assert row[COL["Тело"]] == pytest.approx(float(body["amount"]), abs=TOL), d
            assert row[COL["Остаток после"]] == pytest.approx(float(body["principal_after"]), abs=TOL), d
            checked += 1
        if "sub_payment" in ev:
            assert row[COL["Субсидия к выплате"]] == pytest.approx(float(ev["sub_payment"]["amount"]), abs=TOL), d
            checked += 1
    # в эталоне не осталось платежей, которых нет в Excel (пересмотры ставки в Excel не выводятся)
    leftover = {t for evs in by_date.values() for t in evs if not t.endswith("repricing")}
    assert not leftover
    assert checked > 0


def test_cells_are_formulas():
    wb = openpyxl.load_workbook(OUT)
    for i in XLSX_DEALS:
        ws = wb[f"Сделка {i}"]
        for row in ws.iter_rows(min_row=FIRST_ROW + 1, max_row=FIRST_ROW + 3):
            for name in ("Дата", "Дней", "Доля года τ", "Начислено B·r·τ", "Проценты к выплате", "Тело", "Остаток после"):
                v = row[COL[name]].value
                assert isinstance(v, str) and v.startswith("="), (i, name, v)


def test_file_up_to_date(tmp_path, monkeypatch):
    import make_golden_xlsx

    fresh = tmp_path / "fresh.xlsx"
    monkeypatch.setattr(make_golden_xlsx, "OUT", fresh)
    monkeypatch.setattr(make_golden_xlsx, "ROOT", tmp_path)
    make_golden_xlsx.main()
    a, b = openpyxl.load_workbook(OUT), openpyxl.load_workbook(fresh)
    assert a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        va = [[c.value for c in r] for r in a[name].iter_rows()]
        vb = [[c.value for c in r] for r in b[name].iter_rows()]
        assert va == vb, name


def test_out_path():
    assert OUT == ROOT / "analytics" / "golden_check.xlsx" and Path(OUT).exists()

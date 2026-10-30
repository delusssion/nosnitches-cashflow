#!/usr/bin/env python3
"""Конвертация исходных файлов экспертов в формат контракта (docs/CONTRACT.md).

Берёт xlsx экспертов (testdata/source/) и выдаёт CSV для cfgen:
  deals  — лист «Тестовые сделки» -> файл сделок (формат 1)
  rates  — лист «КС ЦБ»           -> история индикаторов (формат 2)

Официальные файлы в testdata/deals и testdata/rates ведёт аналитик; этот
скрипт — для локальных прогонов и сверки до их появления. Даты в исходнике
встречаются и как DD.MM.YYYY, и как ISO — приводятся к ISO; колонка
«Столбец1» отбрасывается; строки с неизвестным discriminator пропускаются
(их считает валидатор cfgen, не конвертер).

Пример:
  python convert_expert_input.py deals  "../../testdata/source/Расчеты_Analytics_1 сделка.xlsx" deals.csv
  python convert_expert_input.py rates  "../../testdata/source/Расчеты_Analytics_1 сделка.xlsx" key_rate.csv
"""

import argparse
import csv
import sys
from datetime import datetime

# 48 колонок формата 1 — синхронизировано с cfcore::io::DEAL_COLUMNS
DEAL_COLS = [
    "report_dttm", "discriminator", "deal_id", "currency_id",
    "start_principal_amt", "current_principal_amt", "deal_dttm", "maturity_dttm",
    "rate_type_cd", "interest_rate", "base_indicator_rate", "base_indicator",
    "interest_payment_rule_start_dttm", "interest_payment_rule_period_cd",
    "interest_payment_rule_day_num", "day_convention_cd",
    "regular_payment_rule_start_dttm", "regular_payment_rule_period_cd",
    "regular_payment_rule_day_num", "repricing_rule_start_dttm",
    "repricing_rule_period_cd", "repricing_rule_day_num",
    "accrued_interest", "product", "behavior_model_cd",
    "interest_payment", "principal_payment", "repricing_dttm_schedule",
    "repricing_dttm_rate_schedule", "repricing_dttm_rate_schedule_fixed",
    "cap_limit_interest_rate", "floor_limit_interest_rate",
    "sub_rate_type_cd", "sub_interest_rate", "sub_float_indicator",
    "sub_payment_rule_period_cd", "sub_payment_rule_day_num",
    "sub_payment_rule_start_dttm", "sub_repricing_rule_start_dttm",
    "sub_repricing_rule_period_cd", "sub_repricing_rule_day_num",
    "sub_cap_limit_interest_rate", "sub_floor_limit_interest_rate",
    "sub_coeff_index", "deal_report_dttm", "sub_repricing_rule_first_event_flg",
    "indicator_lag_cd", "sub_indicator_lag_cd",
]

# скалярные поля-даты: привести к ISO
DATE_SCALAR = {
    "report_dttm", "deal_dttm", "maturity_dttm",
    "interest_payment_rule_start_dttm", "regular_payment_rule_start_dttm",
    "repricing_rule_start_dttm", "sub_payment_rule_start_dttm",
    "sub_repricing_rule_start_dttm", "deal_report_dttm",
}
# списки дат через пробел
DATE_LIST = {
    "interest_payment", "repricing_dttm_schedule", "repricing_dttm_rate_schedule",
}
# списки пар «дата,сумма» через пробел
PAIR_LIST = {"principal_payment", "repricing_dttm_rate_schedule_fixed"}


def conv_date(v: str) -> str:
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(v, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"не распознана дата: {v!r}")


def conv_date_list(v: str) -> str:
    return " ".join(conv_date(tok) for tok in v.split())


def conv_pairs(v: str) -> str:
    out = []
    for tok in v.split():
        d, amount = tok.rsplit(",", 1)
        out.append(f"{conv_date(d)},{amount}")
    return " ".join(out)


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v).strip()


def convert_deals(src: str, dst: str) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb["Тестовые сделки"]
    it = ws.iter_rows(values_only=True)
    header = [cell_str(h) for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    missing = [c for c in DEAL_COLS if c not in idx]
    if missing:
        sys.exit(f"в листе «Тестовые сделки» нет колонок: {missing}")
    next(it, None)  # строка с описаниями полей

    rows = []
    for r in it:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        disc = r[idx["discriminator"]]
        if disc is None or str(disc).strip() not in {"ANN", "RGM", "SIMPLE"}:
            continue  # служебная строка, не сделка
        out = {}
        for c in DEAL_COLS:
            v = cell_str(r[idx[c]])
            if v:
                if c in DATE_SCALAR:
                    v = conv_date(v)
                elif c in DATE_LIST:
                    v = conv_date_list(v)
                elif c in PAIR_LIST:
                    v = conv_pairs(v)
            out[c] = v
        rows.append(out)

    with open(dst, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=DEAL_COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def convert_rates(src: str, dst: str) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb["КС ЦБ"]
    n = 0
    with open(dst, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["indicator", "date", "rate"])
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            d_raw, r_raw = row[0], row[1]
            if isinstance(d_raw, datetime):
                d = d_raw.date()
            else:
                try:
                    d = datetime.strptime(str(d_raw).strip()[:10], "%Y-%m-%d").date()
                except ValueError:
                    continue  # шапка листа и пустые строки
            try:
                rate = float(str(r_raw).replace(",", "."))
            except (TypeError, ValueError):
                continue
            w.writerow(["CBKEY1M", d.isoformat(), rate])
            n += 1
    return n


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="what", required=True)
    for name in ("deals", "rates"):
        sp = sub.add_parser(name)
        sp.add_argument("src", help="xlsx экспертов (testdata/source/)")
        sp.add_argument("dst", help="куда записать CSV")
    args = p.parse_args()

    n = convert_deals(args.src, args.dst) if args.what == "deals" else convert_rates(args.src, args.dst)
    print(f"{args.what}: {n} строк -> {args.dst}")


if __name__ == "__main__":
    main()
